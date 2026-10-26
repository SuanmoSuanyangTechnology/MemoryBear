"""
Multimodal file processing service.

Processes multimodal files such as images and documents, converting them into formats usable
by LLMs.

Supported providers:
- DashScope (Tongyi Qianwen): OpenAI-compatible content format (image_url / video_url /
  input_audio)
- Bedrock/Anthropic: base64 format only
- OpenAI: URL and base64 formats
"""
import logging

import base64
import csv
import io
import json
import re
import olefile
import struct
import zipfile
from abc import ABC, abstractmethod
from typing import List, Dict, Any, Optional

import PyPDF2
import chardet
import httpx
import magic
import openpyxl
import uuid
from docx import Document
from ipaddress import ip_address
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import Session
from urllib.parse import urlparse

from ..config import current_settings
from ..errors import BizCode
from ..errors import BusinessException
from ..utils.text_sanitize import sanitize_text
from ..models.file_metadata_model import FileMetadata
from ..models.models_model import Modality
from ..schemas.app_schema import FileInput, FileType, FileUploadConfig, TransferMethod
from ..schemas.model_schema import ModelInfo
from ..storage import LocalStorage
from .file_content_service import extract_permanent_file_id
from .file_storage_service import FileStorageService

logger = logging.getLogger(__name__)

TEXT_MIME = ['text/plain', 'text/x-markdown']
PDF_MIME = ['application/pdf']
DOC_MIME = [
    'application/msword',
    'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
]
XLSX_MIME = [
    'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    'application/vnd.ms-excel',
]
CSV_MIME = ['text/csv', 'application/csv']
JSON_MIME = ['application/json']

OPENAI_AUDIO_FORMAT_ALIASES = {
    "mpeg": "mp3",
    "x-mpeg": "mp3",
    "x-wav": "wav",
    "wave": "wav",
    "x-m4a": "m4a",
    "oga": "ogg",
}


def _detect_media_type(content: bytes, declared_type: str | None, fallback: str) -> str:
    """Detect the MIME type from content first, using a controlled fallback when detection fails."""
    try:
        detected_type = magic.from_buffer(content, mime=True)
    except Exception:
        detected_type = None

    if detected_type and detected_type != "application/octet-stream":
        return detected_type.split(";", 1)[0].strip()
    if declared_type:
        return declared_type.split(";", 1)[0].strip()
    return fallback


def _build_data_url(content: bytes, declared_type: str | None, fallback: str) -> str:
    """Build an RFC 2397 Data URL supported by the model API."""
    media_type = _detect_media_type(content, declared_type, fallback)
    encoded_content = base64.b64encode(content).decode("utf-8")
    return f"data:{media_type};base64,{encoded_content}"


def _build_audio_data_url(content: bytes) -> str:
    """Build a Data URL compatible with DashScope/OpenAI audio input."""
    encoded_content = base64.b64encode(content).decode("utf-8")
    return f"data:;base64,{encoded_content}"


MAX_HISTORY_FILE_TEXT_CHARS = 100_000

# Attachments unreachable from the model side must be encoded as Base64/Data URL and sent
# inline. Per-file limits must stay consistent with the app file upload config
# (FileUploadConfig *_max_size_mb per type), otherwise an attachment the user uploaded
# successfully would be silently dropped on the model side; callers that do not pass a config
# explicitly (memory / workflow) simply inherit the FileUploadConfig defaults, keeping a single
# source of truth.
_DEFAULT_UPLOAD_CONFIG = FileUploadConfig()
DEFAULT_INLINE_MAX_SIZE: Dict[str, int] = {
    FileType.IMAGE: _DEFAULT_UPLOAD_CONFIG.image_max_size_mb * 1024 * 1024,
    FileType.AUDIO: _DEFAULT_UPLOAD_CONFIG.audio_max_size_mb * 1024 * 1024,
}
# The per-request inline total is derived from the per-file limits of the types involved in
# this request, avoiding a memory spike when several large attachments are encoded at once; it
# is not exposed as a separate config item.
_FILE_TYPE_LABELS = {
    FileType.IMAGE: "图片",
    FileType.AUDIO: "音频",
    FileType.VIDEO: "视频",
    FileType.DOCUMENT: "文档",
}

# Placeholder shown when video cannot egress: aimed at admins of private deployments,
# keeping the actionable technical clue (unreachable local/intranet storage).
_VIDEO_UNAVAILABLE_HINT = (
    "[视频发送失败：存储地址模型侧不可达（本地存储或内网地址）。"
    "管理员请将对象存储（OSS/S3/MinIO）配置为公网可达后重试，"
    "或接入 provider Files API 支持本地存储视频。]"
)


def _resolve_inline_limits(config: Optional[Dict[str, Any]]) -> Dict[str, int]:
    """Resolve per-type inline single-file limits (raw bytes) from the app file_upload config."""
    limits = dict(DEFAULT_INLINE_MAX_SIZE)
    if isinstance(config, dict):
        mapping = (
            (FileType.IMAGE, "image_max_size_mb"),
            (FileType.AUDIO, "audio_max_size_mb"),
        )
        for file_type, config_key in mapping:
            value = config.get(config_key)
            if value:
                limits[file_type] = int(float(value) * 1024 * 1024)
    return limits


def serialize_file_reference(
        file: FileInput,
        *,
        name: str | None = None,
        size: int | None = None,
) -> Dict[str, Any]:
    """Serialize an attachment into a reference that is safe to persist and rehydrate."""
    is_local = file.transfer_method == TransferMethod.LOCAL_FILE
    return {
        "type": file.type,
        "transfer_method": file.transfer_method.value,
        "upload_file_id": str(file.upload_file_id) if is_local and file.upload_file_id else None,
        # Private URLs must not be written into message metadata; local files must be re-read
        # by ID after a fresh ownership check.
        "url": None if is_local else file.url,
        "name": name if name is not None else file.name,
        "size": size if size is not None else file.size,
        "file_type": file.file_type,
    }


def _built_permanent_file_id(url: Any, *, any_host: bool = False) -> str | None:
    """Recognize a permanent download URL minted by this service and extract the file ID from it.

    The implementation was extracted into the shared file content service
    (file_content_service.extract_permanent_file_id) and is shared with consumers such as
    knowledge base image retrieval, so the "does this URL belong to this service" decision does
    not slowly fork per call site in private deployments.
    """
    return extract_permanent_file_id(url, trust_any_host=any_host)


def _model_cannot_reach_url(url: Any) -> bool:
    """Decide whether the model server can reach this URL.

    The decision uses the hostname / literal IP only, with no DNS resolution: a missing
    hostname, localhost, loopback, private ranges, link-local/reserved addresses, and this
    platform's own domain (FILE_LOCAL_SERVER_URL pointing at localhost or an intranet) all
    count as unreachable from the model, in which case the attachment must egress as inline
    bytes instead. Public domains or public IPs count as reachable, preserving native URL
    pass-through.
    """
    if not isinstance(url, str) or not url.strip():
        return True

    host = urlparse(url.strip()).hostname
    if not host:
        return True

    host = host.strip("[]").lower()
    if host == "localhost" or host.endswith(".localhost"):
        return True

    own_host = urlparse(current_settings().file_local_server_url or "").hostname
    if own_host and host == own_host.strip("[]").lower():
        return True

    try:
        parsed_ip = ip_address(host)
    except ValueError:
        # Domains are not resolved: the result varies with the deployment environment, so
        # public domains count as reachable.
        return False

    return bool(
        parsed_ip.is_loopback
        or parsed_ip.is_private
        or parsed_ip.is_link_local
        or parsed_ip.is_reserved
        or parsed_ip.is_unspecified
    )


def _requires_inline_bearer(file: FileInput) -> bool:
    """Decide whether this attachment must egress with bytes read by this service.

    Any of these conditions forces inlining:
    1. an explicitly local_file attachment;
    2. the URL is a permanent download address minted by this service (whatever
       transfer_method says);
    3. the URL is unreachable from the model side (intranet/loopback) and a file ID allows
       fetching the bytes directly, bypassing HTTP.
    """
    if file.transfer_method == TransferMethod.LOCAL_FILE:
        return True
    if not file.url:
        return False
    if _built_permanent_file_id(file.url):
        return True
    return bool(file.upload_file_id) and _model_cannot_reach_url(file.url)


def deserialize_file_reference(reference: Dict[str, Any]) -> FileInput:
    """Restore FileInput from current or legacy message metadata, without using a private URL as a remote file."""
    url = reference.get("url")
    upload_file_id = reference.get("upload_file_id")
    transfer_method_value = reference.get("transfer_method")

    try:
        transfer_method = TransferMethod(transfer_method_value) if transfer_method_value else None
    except ValueError:
        transfer_method = None

    if transfer_method is None:
        if upload_file_id:
            transfer_method = TransferMethod.LOCAL_FILE
        else:
            legacy_file_id = _built_permanent_file_id(url, any_host=True)
            if legacy_file_id:
                upload_file_id = legacy_file_id
                transfer_method = TransferMethod.LOCAL_FILE
            else:
                transfer_method = TransferMethod.REMOTE_URL

    payload: Dict[str, Any] = {
        "type": reference.get("type", FileType.DOCUMENT),
        "transfer_method": transfer_method,
        "file_type": reference.get("file_type"),
        "name": reference.get("name"),
        "size": reference.get("size"),
    }
    if transfer_method == TransferMethod.LOCAL_FILE:
        payload["upload_file_id"] = upload_file_id
    else:
        payload["url"] = url
    return FileInput(**payload)


def sanitize_processed_files_for_history(processed_files: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Strip inline binary content, keeping bounded text or an attachment-presence notice."""
    def is_data_url(value: Any) -> bool:
        return isinstance(value, str) and value.startswith("data:")

    def binary_history_projection(part: Dict[str, Any]) -> Dict[str, Any]:
        source = part.get("source")
        if isinstance(source, dict) and source.get("type") == "base64" and source.get("media_type") == "text/plain":
            try:
                text = base64.b64decode(source.get("data", "")).decode("utf-8", errors="replace")
                return {"type": "text", "text": text[:MAX_HISTORY_FILE_TEXT_CHARS]}
            except (TypeError, ValueError):
                pass

        part_type = str(part.get("type") or "附件")
        return {
            "type": "text",
            "text": f"[{part_type} 附件已在原始请求中处理；二进制内容未保存到会话历史]",
        }

    sanitized_files: List[Dict[str, Any]] = []
    inline_binary_count = 0
    for part in processed_files:
        if not isinstance(part, dict):
            continue

        source = part.get("source")
        input_audio = part.get("input_audio")
        image_url = part.get("image_url")
        audio_url = part.get("audio_url")
        video_url = part.get("video_url")
        contains_inline_binary = (
            isinstance(source, dict) and source.get("type") == "base64"
        ) or (
            isinstance(input_audio, dict) and is_data_url(input_audio.get("data"))
        ) or (
            isinstance(image_url, dict) and is_data_url(image_url.get("url"))
        ) or (
            isinstance(audio_url, dict) and is_data_url(audio_url.get("url"))
        ) or (
            isinstance(video_url, dict) and is_data_url(video_url.get("url"))
        ) or is_data_url(part.get("image")) or is_data_url(part.get("audio")) or is_data_url(part.get("video"))

        if contains_inline_binary:
            inline_binary_count += 1
            sanitized_files.append(binary_history_projection(part))
        else:
            sanitized_files.append(part)

    if inline_binary_count:
        logger.info("已从会话历史移除 %s 个内联附件内容", inline_binary_count)
    return sanitized_files


class MultimodalFormatStrategy(ABC):
    """Multimodal format strategy base class."""

    def __init__(self, file: FileInput):
        self.file = file

    @abstractmethod
    async def format_image(self, url: str | None, content: bytes | None = None) -> tuple[bool, Dict[str, Any]]:
        """Format an image."""
        pass

    @abstractmethod
    async def format_document(self, file_name: str, text: str) -> tuple[bool, Dict[str, Any]]:
        """Format a document."""
        pass

    @abstractmethod
    async def format_audio(
            self,
            file_type: str,
            url: str | None,
            content: bytes | None = None,
    ) -> tuple[bool, Dict[str, Any]]:
        """Format audio."""
        pass

    @abstractmethod
    async def format_video(self, url: str | None, content: bytes | None = None) -> tuple[bool, Dict[str, Any]]:
        """Format a video."""
        pass


class BedrockFormatStrategy(MultimodalFormatStrategy):
    """Bedrock/Anthropic strategy."""

    async def format_image(self, url: str | None, content: bytes | None = None) -> tuple[bool, Dict[str, Any]]:
        """
        Bedrock/Anthropic format: base64 encoding
        {"type": "image", "source": {"type": "base64", "media_type": "...", "data": "..."}}
        """
        if content is None:
            if not url:
                return False, {"type": "text", "text": "[图片文件缺少可用内容]"}
            logger.info(f"下载并编码图片: {url}")
            async with httpx.AsyncClient(timeout=30.0) as client:
                response = await client.get(url, follow_redirects=True)
                response.raise_for_status()
                content = response.content
                self.file.set_content(content)

        media_type = _detect_media_type(content, self.file.file_type, "image/jpeg")
        if not media_type.startswith("image/"):
            media_type = "image/jpeg"
        base64_data = base64.b64encode(content).decode("utf-8")

        logger.info(f"图片编码完成: media_type={media_type}, size={len(base64_data)}")

        return True, {
            "type": "image",
            "source": {
                "type": "base64",
                "media_type": media_type,
                "data": base64_data,
            },
        }

    async def format_document(self, file_name: str, text: str) -> tuple[bool, Dict[str, Any]]:
        """Bedrock/Anthropic document format (base64 encoding required)."""
        text = f"文档内容：\n{text}\n"
        text_bytes = text.encode("utf-8")
        base64_text = base64.b64encode(text_bytes).decode("utf-8")

        return True, {
            "type": "document",
            "source": {
                "type": "base64",
                "media_type": "text/plain",
                "data": base64_text,
            },
        }

    async def format_audio(
            self,
            file_type: str,
            url: str | None,
            content: bytes | None = None,
    ) -> tuple[bool, Dict[str, Any]]:
        """Bedrock/Anthropic does not support native audio."""
        return False, {
            "type": "text",
            "text": "[音频文件：当前 provider 不支持原生音频]",
        }

    async def format_video(self, url: str | None, content: bytes | None = None) -> tuple[bool, Dict[str, Any]]:
        """Bedrock/Anthropic video format."""
        return False, {
            "type": "text",
            "text": "[视频文件：当前 provider 暂不支持]",
        }


class OpenAIFormatStrategy(MultimodalFormatStrategy):
    """OpenAI-compatible strategy."""

    async def format_image(self, url: str | None, content: bytes | None = None) -> tuple[bool, Dict[str, Any]]:
        """Local images use a Data URL, so the model side never pulls intranet URLs."""
        image_source = _build_data_url(content, self.file.file_type, "image/jpeg") if content is not None else url
        if not image_source:
            return False, {"type": "text", "text": "[图片文件缺少可用内容]"}
        return True, {
            "type": "image_url",
            "image_url": {
                "url": image_source,
            },
        }

    async def format_document(self, file_name: str, text: str) -> tuple[bool, Dict[str, Any]]:
        """OpenAI document format."""
        return True, {
            "type": "text",
            "text": f"<document name=\"{file_name}\">\n文档内容：\n{text}\n</document>",
        }

    async def format_audio(
            self,
            file_type: str,
            url: str | None,
            content: bytes | None = None,
    ) -> tuple[bool, Dict[str, Any]]:
        """OpenAI audio format."""
        try:
            audio_data = content
            if audio_data is None:
                if not url:
                    raise ValueError("音频文件缺少可用内容")
                async with httpx.AsyncClient(timeout=30.0) as client:
                    response = await client.get(url, follow_redirects=True)
                    response.raise_for_status()
                    audio_data = response.content
                    self.file.set_content(audio_data)

            file_ext = file_type.split("/")[-1] if file_type and "/" in file_type else None
            if not file_ext:
                content_type = _detect_media_type(audio_data, self.file.file_type, "audio/wav")
                file_ext = content_type.split("/")[-1].split(";", 1)[0].strip() if "/" in content_type else None
            if not file_ext and url:
                file_ext = url.split("?", 1)[0].rsplit(".", 1)[-1].lower() or None
            file_ext = "wav" if not file_ext else file_ext.lower()
            file_ext = OPENAI_AUDIO_FORMAT_ALIASES.get(file_ext, file_ext)

            return True, {
                "type": "input_audio",
                "input_audio": {
                    "data": _build_audio_data_url(audio_data),
                    "format": file_ext,
                },
            }
        except Exception as e:
            logger.error(f"下载音频失败: {e}")
            return False, {
                "type": "text",
                "text": f"[音频处理失败: {str(e)}]",
            }

    async def format_video(self, url: str | None, content: bytes | None = None) -> tuple[bool, Dict[str, Any]]:
        """OpenAI video format; local videos need the provider Files API to be wired up later."""
        if content is not None:
            return False, {
                "type": "text",
                "text": _VIDEO_UNAVAILABLE_HINT,
            }
        if not url:
            return False, {"type": "text", "text": "[视频文件缺少可用 URL]"}
        return True, {
            "type": "video_url",
            "video_url": {
                "url": url,
            },
        }


# Provider-to-strategy mapping
PROVIDER_STRATEGIES = {
    "dashscope": OpenAIFormatStrategy,
    "bedrock": BedrockFormatStrategy,
    "anthropic": BedrockFormatStrategy,
    "openai": OpenAIFormatStrategy,
    "volcano": OpenAIFormatStrategy,
    "speedbear": OpenAIFormatStrategy,
    "gpustack": OpenAIFormatStrategy,
}


class MultimodalService:
    """
    Service for handling multimodal file processing.

    Attributes:
        db (Session): Database session.
        model_view (ModelInfo | None): non-decrypted runtime model view (no longer holds credentials).
        provider (str): Name of the model provider.
        input_modalities (list): Input modality configuration of the model.
    """

    def __init__(
            self,
            db: Session | AsyncSession,
            model_view: ModelInfo | None = None,
    ):
        """
        Initialize the multimodal service.

        Args:
            db (Session): Database session.
            model_view (ModelInfo | None): non-decrypted runtime model view (credentials live on the model service side).
        """
        self.db = db
        self.model_view = model_view
        if self.model_view is not None:
            self.provider = model_view.provider.lower()
            self.input_modalities = list(getattr(model_view, "input_modalities", None) or [])

    def _uses_async_session(self) -> bool:
        return isinstance(self.db, AsyncSession)

    async def _get_file_metadata(
            self,
            file_id: uuid.UUID | None,
            *,
            completed_only: bool = False,
    ) -> FileMetadata | None:
        if file_id is None:
            return None

        if self._uses_async_session():
            stmt = select(FileMetadata).where(FileMetadata.id == file_id)
            if completed_only:
                stmt = stmt.where(FileMetadata.status == "completed")
            return (await self.db.execute(stmt)).scalar_one_or_none()

        query = self.db.query(FileMetadata).filter(FileMetadata.id == file_id)
        if completed_only:
            query = query.filter(FileMetadata.status == "completed")
        return query.first()

    async def _get_local_file_metadata(
            self,
            file_id: uuid.UUID | None,
            workspace_id: uuid.UUID | str | None,
    ) -> FileMetadata:
        """Fetch a model-processable local attachment scoped to the current workspace or tenant."""
        if file_id is None:
            raise BusinessException("本地文件 ID 不能为空", BizCode.FILE_NOT_FOUND)
        if self.db is None:
            raise BusinessException("缺少读取本地文件所需的数据库上下文", BizCode.FILE_READ_ERROR)

        tenant_id = getattr(self.model_view, "tenant_id", None) if self.model_view else None
        if workspace_id is None and tenant_id is None:
            raise BusinessException(
                "缺少工作空间或租户上下文，无法读取本地文件",
                BizCode.FILE_NOT_FOUND,
            )

        if self._uses_async_session():
            stmt = select(FileMetadata).where(
                FileMetadata.id == file_id,
                FileMetadata.status == "completed",
            )
            if workspace_id is not None:
                stmt = stmt.where(FileMetadata.workspace_id == workspace_id)
            if tenant_id is not None:
                stmt = stmt.where(FileMetadata.tenant_id == tenant_id)
            metadata = (await self.db.execute(stmt)).scalar_one_or_none()
        else:
            query = self.db.query(FileMetadata).filter(
                FileMetadata.id == file_id,
                FileMetadata.status == "completed",
            )
            if workspace_id is not None:
                query = query.filter(FileMetadata.workspace_id == workspace_id)
            if tenant_id is not None:
                query = query.filter(FileMetadata.tenant_id == tenant_id)
            metadata = query.first()

        if metadata is None:
            # Return unavailable uniformly to callers, so other workspaces' file existence is
            # not leaked.
            raise BusinessException("文件不存在、未完成或无权访问", BizCode.FILE_NOT_FOUND)
        return metadata

    async def _resolve_video_public_url(
            self,
            file: FileInput,
            workspace_id: uuid.UUID | str | None,
    ) -> str | None:
        """Mint a model-reachable direct URL for an uploaded video.

        Video has no universal inline (base64) protocol, so remote storage backends
        (OSS/S3/MinIO) issue a long-lived presigned object URL instead of this service's
        own proxy address. The minted URL still has to pass the same model-side
        reachability check; intranet MinIO endpoints and Aliyun VPC-internal endpoints
        are rejected here as well. Returns None for local disk storage or unreachable
        backends, leaving the caller to emit the existing error placeholder.
        """
        if not file.upload_file_id:
            return None

        storage = FileStorageService().storage
        if isinstance(storage, LocalStorage):
            logger.info(
                "视频无法出站：当前为本地磁盘存储，无模型侧可达地址",
                extra={"file_id": str(file.upload_file_id)},
            )
            return None

        # Ownership/tenant check before minting a signed URL.
        metadata = await self._get_local_file_metadata(file.upload_file_id, workspace_id)
        try:
            # No file_name: do not force Content-Disposition: attachment, which some
            # model-side video fetchers reject.
            url = await storage.get_url(metadata.file_key, expires=86400)
        except Exception as exc:
            logger.warning(f"视频预签名 URL 生成失败，已降级: {exc}")
            return None

        host = (urlparse(url).hostname or "").strip("[]").lower()
        if _model_cannot_reach_url(url) or "-internal." in host:
            logger.warning(
                "视频存储地址模型侧不可达，已拦截",
                extra={"file_id": str(metadata.id), "host": host},
            )
            return None
        return url

    async def _hydrate_local_file(
            self,
            file: FileInput,
            workspace_id: uuid.UUID | str | None,
    ) -> int:
        """Read bytes of an attachment unreachable from the model side, for inline use by the
        egress adapter.

        - With an upload_file_id: a workspace/tenant-scoped query + direct storage read, which
          also backfills file name, size and MIME for downstream consumers such as perception
          memory;
        - Otherwise only permanent URLs minted by this service are recognized and fetched over
          HTTP by this service itself; the URL is never handed to the model, and arbitrary
          remote URLs are not accepted (avoiding SSRF).
        """
        existing_content = file.get_content()
        if existing_content is not None:
            if len(existing_content) > current_settings().max_file_size:
                raise BusinessException("文件大小不在允许范围内", BizCode.FILE_READ_ERROR)
            return len(existing_content)

        if file.upload_file_id:
            metadata = await self._get_local_file_metadata(file.upload_file_id, workspace_id)
            file_size = int(metadata.file_size or 0)
            if file_size <= 0 or file_size > current_settings().max_file_size:
                raise BusinessException("文件大小不在允许范围内", BizCode.FILE_READ_ERROR)

            content = await FileStorageService().download_file(metadata.file_key)
            if len(content) != file_size:
                logger.error(
                    "本地文件大小与元数据不一致",
                    extra={"file_id": str(metadata.id), "expected_size": file_size, "actual_size": len(content)},
                )
                raise BusinessException("文件读取失败", BizCode.FILE_READ_ERROR)

            file.set_content(content)
            file.name = metadata.file_name
            file.size = file_size
            file.file_type = _detect_media_type(
                content,
                metadata.content_type or file.file_type,
                "application/octet-stream",
            )
            return file_size

        if not _built_permanent_file_id(file.url):
            # Unreachable URL not minted by this service: do not self-fetch; let the strategy
            # layer degrade, avoiding SSRF.
            raise BusinessException("附件 URL 模型侧不可达且无法定位文件", BizCode.FILE_READ_ERROR)

        try:
            async with httpx.AsyncClient(timeout=30.0) as client:
                response = await client.get(file.url, follow_redirects=True)
                response.raise_for_status()
                content = response.content
        except Exception as exc:
            logger.error(f"读取本服务附件 URL 失败: {exc}", exc_info=True)
            raise BusinessException("文件读取失败", BizCode.FILE_READ_ERROR)

        if len(content) > current_settings().max_file_size:
            raise BusinessException("文件大小不在允许范围内", BizCode.FILE_READ_ERROR)

        file.set_content(content)
        return len(content)

    async def process_files(
            self,
            files: Optional[List[FileInput]],
            workspace_id: uuid.UUID = None,
            document_image_recognition: bool = False,
            include_processing_errors: bool = True,
            file_upload_config: Optional[Dict[str, Any]] = None,
    ) -> List[Dict[str, Any]]:
        """
        Process a file list and return LLM-usable formats.

        Args:
            files: File input list.
            include_processing_errors: Whether to append placeholder text for failed files to
                the result; when False, failed files are skipped and later files keep being
                processed.
            file_upload_config: App file upload config (FileUploadConfig.model_dump); inline
                size limits follow its per-type *_max_size_mb values.

        Returns:
            List[Dict]: List of LLM-usable content formats (format varies by provider).
        """
        if not files:
            return []

        # The per-file inline limit comes from the app config, ensuring "uploadable means
        # sendable to the model request"; the total is derived as 2x the largest per-file limit
        # among the types involved (consistent with the historical criteria) to cap the memory
        # peak when multiple attachments are encoded at once.
        inline_limits = _resolve_inline_limits(file_upload_config)
        present_limits = [
            inline_limits[f.type] for f in files if f.type in inline_limits
        ]
        inline_total_limit = max(present_limits, default=0) * 2

        # Resolve the strategy: providers that have unified on the OpenAI-compatible protocol
        # (dashscope/volcano/openai/minimax/...) must all produce OpenAI multimodal format.
        # is_omni no longer participates in the decision -- after the native ChatTongyi
        # protocol was retired, a native {"type": "image", ...} is only rejected by the
        # compatible endpoint with 400 invalid_value (see
        # core/models/base.py:get_provider_llm_class).
        strategy_class = PROVIDER_STRATEGIES.get(self.provider)
        if not strategy_class:
            logger.warning(
                f"未找到 provider '{self.provider}' 的策略，使用 OpenAI 兼容格式"
            )
            strategy_class = OpenAIFormatStrategy

        result = []
        inline_payload_size = 0
        for idx, file in enumerate(files):
            # Attachments unreachable from the model side can only egress after this service
            # reads their bytes; FILE_LOCAL_SERVER_URL or a private OSS URL must never be
            # handed to the model to pull back. Publicly reachable URLs are passed through
            # as-is, with no extra byte shuttling.
            if _requires_inline_bearer(file):
                if file.type == FileType.VIDEO:
                    # Video has no universal inline protocol: remote storage (OSS/S3/MinIO)
                    # hands the model a direct, model-reachable presigned URL. Only local
                    # disk storage or unreachable intranet endpoints keep the explicit error.
                    video_public_url = None
                    if Modality.VIDEO in self.input_modalities:
                        # Ownership/tenant failure inside still aborts the request as before.
                        video_public_url = await self._resolve_video_public_url(file, workspace_id)

                    if video_public_url is not None:
                        file.url = video_public_url
                        is_support, content = await self._process_video(file, strategy_class(file))
                        if is_support or include_processing_errors:
                            result.append(content)
                    elif Modality.VIDEO in self.input_modalities and include_processing_errors:
                        result.append({
                            "type": "text",
                            "text": _VIDEO_UNAVAILABLE_HINT,
                        })
                    elif Modality.VIDEO not in self.input_modalities:
                        logger.warning(f"不支持的文件类型: {file.type}")
                    continue

                try:
                    inline_size = await self._hydrate_local_file(file, workspace_id)
                except BusinessException as exc:
                    # A single attachment that cannot be inlined degrades only itself, not the
                    # other files in the same request.
                    logger.warning(f"附件内联读取失败，已降级: {exc}")
                    if include_processing_errors:
                        result.append({"type": "text", "text": f"[文件处理失败: {str(exc)}]"})
                    continue
                if file.type in {FileType.IMAGE, FileType.AUDIO}:
                    file_type_limit = inline_limits.get(file.type)
                    if file_type_limit is None:
                        # Types without a configured inline limit get no size cap (e.g.
                        # future new types).
                        inline_payload_size += inline_size
                    elif inline_size > file_type_limit:
                        logger.warning(
                            "附件超过应用配置的内联大小限制",
                            extra={"file_id": str(file.upload_file_id), "size": inline_size},
                        )
                        if include_processing_errors:
                            limit_mb = file_type_limit // (1024 * 1024)
                            result.append({
                                "type": "text",
                                "text": f"[{_FILE_TYPE_LABELS.get(file.type, '文件')}超过单文件大小限制 {limit_mb}MB，未发送给模型，请压缩或拆分后重试]",
                            })
                        continue
                    elif inline_payload_size + inline_size > inline_total_limit:
                        logger.warning(
                            "本次请求内联附件总大小超过限制",
                            extra={"file_id": str(file.upload_file_id), "size": inline_size},
                        )
                        if include_processing_errors:
                            result.append({
                                "type": "text",
                                "text": "[本次请求附件总大小超限，该文件未发送给模型，请减少附件后重试]",
                            })
                        continue
                    else:
                        inline_payload_size += inline_size
            elif not file.url:
                file.url = await self.get_file_url(file)

            strategy = strategy_class(file)
            try:
                if file.type == FileType.IMAGE and Modality.IMAGE in self.input_modalities:
                    is_support, content = await self._process_image(file, strategy)
                    if is_support or include_processing_errors:
                        result.append(content)
                elif file.type == FileType.DOCUMENT:
                    is_support, content = await self._process_document(file, strategy)
                    if not is_support and not include_processing_errors:
                        # Skip only the current failed document; later files keep being
                        # processed. If nothing succeeds at all, the caller falls back to the
                        # existing summary.
                        continue
                    result.append(content)
                    # Extract embedded document images only when the switch is on and the
                    # model supports vision
                    if document_image_recognition and Modality.IMAGE in self.input_modalities:
                        img_infos = await self.extract_document_images(file)
                        img_result = []
                        for img_info in img_infos:
                            page = img_info["page"]
                            index = img_info["index"]
                            ext = img_info.get("ext", "png")
                            image_bytes = img_info["bytes"]
                            if len(image_bytes) > inline_limits.get(FileType.IMAGE, 0) or (
                                    inline_payload_size + len(image_bytes) > inline_total_limit
                            ):
                                logger.warning(
                                    "文档内嵌图片超过内联大小限制，已跳过",
                                    extra={"page": page, "index": index, "size": len(image_bytes)},
                                )
                                continue
                            inline_payload_size += len(image_bytes)
                            try:
                                placeholder = f"第{page}页 第{index + 1}张" if page > 0 else f"第{index + 1}张"
                                # The text keeps only positional markers; display URLs belong to
                                # the frontend/storage layer and must not enter the public-model
                                # context.
                                if result and result[-1].get("type") in ("text", "document"):
                                    key = "text" if "text" in result[-1] else list(result[-1].keys())[-1]
                                    result[-1][key] = result[-1].get(key, "") + f"\n[图片 {placeholder}]"
                                # Embedded images are given to the vision model directly as
                                # bytes; no intranet permanent URL is created or exposed.
                                img_file = FileInput(
                                    type=FileType.IMAGE,
                                    transfer_method=TransferMethod.LOCAL_FILE,
                                    upload_file_id=file.upload_file_id or uuid.uuid4(),
                                    file_type=f"image/{ext}",
                                )
                                img_file.set_content(image_bytes)
                                img_support, img_content = await self._process_image(img_file, strategy_class(img_file))
                                if img_support or include_processing_errors:
                                    img_result.append(img_content)
                            except Exception as img_err:
                                logger.warning(f"文档图片处理失败: {img_err}")
                        result.extend(img_result)
                elif file.type == FileType.AUDIO and Modality.AUDIO in self.input_modalities:
                    is_support, content = await self._process_audio(file, strategy)
                    if is_support or include_processing_errors:
                        result.append(content)
                elif file.type == FileType.VIDEO and Modality.VIDEO in self.input_modalities:
                    is_support, content = await self._process_video(file, strategy)
                    if is_support or include_processing_errors:
                        result.append(content)
                else:
                    logger.warning(f"不支持的文件类型: {file.type}")
            except Exception as e:
                logger.error(
                    f"处理文件失败",
                    extra={
                        "file_index": idx,
                        "file_type": file.type,
                        "error": str(e)
                    },
                    exc_info=True
                )
                # By default an error placeholder is kept for history; when disabled, only
                # the current failed file is skipped and later files keep being processed.
                if include_processing_errors:
                    result.append({
                        "type": "text",
                        "text": f"[文件处理失败: {str(e)}]"
                    })

        logger.info(f"成功处理 {len(result)}/{len(files)} 个文件，provider={self.provider}")
        return result

    async def _process_image(self, file: FileInput, strategy) -> tuple[bool, Dict[str, Any]]:
        """Process an image file."""
        try:
            return await strategy.format_image(file.url, content=file.get_content())
        except Exception as e:
            logger.error(f"处理图片失败: {e}", exc_info=True)
            return False, {
                "type": "text",
                "text": f"[图片处理失败: {str(e)}]",
            }

    async def _process_document(self, file: FileInput, strategy) -> tuple[bool, Dict[str, Any]]:
        """Process a document file (PDF, Word, etc.); only extracted content is passed to the model."""
        if not _requires_inline_bearer(file):
            return True, {
                "type": "text",
                "text": f"<document url=\"{file.url}\">\n{await self.extract_document_text(file)}\n</document>",
            }

        text = await self.extract_document_text(file)
        return await strategy.format_document(file.name or "unknown", text)

    async def _process_audio(self, file: FileInput, strategy) -> tuple[bool, Dict[str, Any]]:
        """Process an audio file (native audio inline; transcription belongs to the model service asr family)."""
        try:
            return await strategy.format_audio(
                file.file_type or "",
                file.url,
                file.get_content(),
            )
        except Exception as e:
            logger.error(f"处理音频失败: {e}", exc_info=True)
            return False, {
                "type": "text",
                "text": f"[音频处理失败: {str(e)}]",
            }

    async def _process_video(self, file: FileInput, strategy) -> tuple[bool, Dict[str, Any]]:
        """Process a video file."""
        try:
            return await strategy.format_video(file.url, content=file.get_content())
        except Exception as e:
            logger.error(f"处理视频失败: {e}", exc_info=True)
            return False, {
                "type": "text",
                "text": f"[视频处理失败: {str(e)}]",
            }

    async def get_file_url(self, file: FileInput) -> str:
        """
        Get a file's access URL.
        
        Args:
            file: File Input Struct
            
        Returns:
            str: File access URL (permanent URL).
            
        Raises:
            BusinessException: file does not exist.
        """
        if file.transfer_method == TransferMethod.REMOTE_URL:
            return file.url
        else:
            file_id = file.upload_file_id

            # Query FileMetadata
            file_metadata = await self._get_file_metadata(file_id, completed_only=True)

            if not file_metadata:
                raise BusinessException(
                    f"文件不存在或已删除: {file_id}",
                    BizCode.NOT_FOUND
                )

            # Return the permanent URL
            server_url = current_settings().file_local_server_url
            return f"{server_url}/storage/permanent/{file_id}"

    async def extract_document_text(self, file: FileInput) -> str:
        """
        Extract document text content.
        
        Args:
            file: File input.
            
        Returns:
            str: Extracted text content.
        """
        try:
            file_content = file.get_content()
            if not file_content:
                async with httpx.AsyncClient(timeout=30.0) as client:
                    response = await client.get(file.url, follow_redirects=True)
                    response.raise_for_status()
                    file_content = response.content
                    file.set_content(file_content)
            file_mime_type = magic.from_buffer(file_content, mime=True)
            if file_mime_type in TEXT_MIME:
                return self._decode_text_safe(file_content)
            elif file_mime_type in PDF_MIME:
                return await self._extract_pdf_text(file_content)
            elif self._is_word_file(file_content, file_mime_type):
                return await self._extract_word_text(file_content)
            elif self._is_excel_file(file_content, file_mime_type):
                return await self._extract_xlsx_text(file_content)
            elif file_mime_type in CSV_MIME:
                return await self._extract_csv_text(file_content)
            elif file_mime_type in JSON_MIME:
                return await self._extract_json_text(file_content)
            else:
                return f"[Unsupported file type: {file_mime_type}]"
        except Exception as e:
            logger.error(f"Failed to load file. - {e}")
            return "[Failed to load file.]"

    async def extract_document_images(self, file: FileInput) -> list[dict]:
        """
        Extract embedded images from a document (PDF and DOCX supported) with position info.

        Returns:
            list[dict]: Each item contains:
                - bytes: image binary
                - page: page number (PDF starts at 1; DOCX uses 0)
                - index: image ordinal within the page/document (0-based)
                - ext: image extension (e.g. png, jpeg)
        """
        try:
            file_content = file.get_content()
            if not file_content:
                async with httpx.AsyncClient(timeout=30.0) as client:
                    response = await client.get(file.url, follow_redirects=True)
                    response.raise_for_status()
                    file_content = response.content
                    file.set_content(file_content)

            file_mime_type = magic.from_buffer(file_content, mime=True)
            if file_mime_type in PDF_MIME:
                return self._extract_pdf_images(file_content)
            elif self._is_word_file(file_content, file_mime_type):
                return self._extract_docx_images(file_content)
            return []
        except Exception as e:
            logger.error(f"提取文档图片失败: {e}")
            return []

    @staticmethod
    def _extract_pdf_images(file_content: bytes) -> list[dict]:
        """Extract embedded images from a PDF with page number and ordinal."""
        images = []
        import fitz  # PyMuPDF
        try:
            doc = fitz.open(stream=file_content, filetype="pdf")
            for page_num, page in enumerate(doc, start=1):
                for idx, img in enumerate(page.get_images(full=True)):
                    xref = img[0]
                    base_image = doc.extract_image(xref)
                    images.append({
                        "bytes": base_image["image"],
                        "ext": base_image.get("ext", "png"),
                        "page": page_num,
                        "index": idx,
                    })
            doc.close()
        except ImportError:
            logger.warning("PyMuPDF 未安装，无法提取 PDF 图片，请执行: uv add pymupdf")
        except Exception as e:
            logger.error(f"提取 PDF 图片失败: {e}")
        return images

    @staticmethod
    def _extract_docx_images(file_content: bytes) -> list[dict]:
        """Extract embedded images from a DOCX with ordinal (DOCX has no page numbers; page is fixed at 0)."""
        images = []
        try:
            if file_content[:2] != b'PK':
                return []
            with zipfile.ZipFile(io.BytesIO(file_content)) as zf:
                media_files = sorted(
                    name for name in zf.namelist()
                    if name.startswith("word/media/") and not name.endswith("/")
                )
                for idx, name in enumerate(media_files):
                    ext = name.rsplit(".", 1)[-1].lower() if "." in name else "png"
                    images.append({
                        "bytes": zf.read(name),
                        "ext": ext,
                        "page": 0,
                        "index": idx,
                    })
        except Exception as e:
            logger.error(f"提取 DOCX 图片失败: {e}")
        return images

    @staticmethod
    async def _extract_pdf_text(file_content: bytes) -> str:
        """Extract PDF text."""
        try:
            # Read the PDF via BytesIO
            text_parts = []
            pdf_file = io.BytesIO(file_content)
            pdf_reader = PyPDF2.PdfReader(pdf_file)
            for page in pdf_reader.pages:
                text_parts.append(page.extract_text())
            return '\n'.join(text_parts)
        except Exception as e:
            logger.error(f"提取 PDF 文本失败: {e}")
            return f"[PDF 提取失败: {str(e)}]"

    @staticmethod
    async def _extract_word_text(file_content: bytes) -> str:
        """Extract text from a Word document (supports .docx and legacy .doc)."""
        # Try docx first (ZIP format)
        if file_content[:2] == b'PK':
            try:
                word_file = io.BytesIO(file_content)
                doc = Document(word_file)
                text_lines = []
                for p in doc.paragraphs:
                    text = p.text.strip()
                    if text:
                        text_lines.append(text)

                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            text = cell.text.strip()
                            if text:
                                text_lines.append(text)

                full_text = "\n".join(text_lines)
                return full_text.strip() or "[docx 文件无文本内容]"
            except Exception as e:
                logger.error(f"提取 docx 文本失败: {str(e)}", exc_info=True)
                return f"[docx 提取失败: {str(e)}]"

        # Legacy .doc (OLE2/CFB format); parse the piece table per the Word Binary Format
        try:
            ole = olefile.OleFileIO(io.BytesIO(file_content))
            word_stream = ole.openstream('WordDocument').read()

            # FIB offset 0xA bit9 chooses between 0Table and 1Table
            fib_flags = struct.unpack_from('<H', word_stream, 0xA)[0]
            table_name = '1Table' if (fib_flags & 0x0200) else '0Table'
            table_stream = ole.openstream(table_name).read()

            # Read fcClx/lcbClx from the FIB to locate the piece table
            fc_clx, lcb_clx = struct.unpack_from("<II", word_stream, 0x1A2)
            clx = table_stream[fc_clx: fc_clx + lcb_clx]

            # Parse the CLX to find the PlcPcd (piece table)
            i, plc_pcd = 0, None
            while i < len(clx):
                clxt = clx[i]
                if clxt == 0x01:
                    i += 3 + struct.unpack_from('<H', clx, i + 1)[0]
                elif clxt == 0x02:
                    cb = struct.unpack_from('<I', clx, i + 1)[0]
                    plc_pcd = clx[i + 5: i + 5 + cb]
                    break
                else:
                    break

            if plc_pcd is None:
                raise ValueError("PlcPcd not found")

            # PlcPcd: (n+1) CP entries (4 bytes each) + n PCD entries (8 bytes each)
            n_pieces = (len(plc_pcd) - 4) // 12
            cp_array = [struct.unpack_from('<I', plc_pcd, k * 4)[0] for k in range(n_pieces + 1)]

            parts = []
            for k in range(n_pieces):
                fc_value = struct.unpack_from('<I', plc_pcd, (n_pieces + 1) * 4 + k * 8 + 2)[0]
                is_ansi = bool(fc_value & 0x40000000)
                fc = fc_value & 0x3FFFFFFF
                char_count = cp_array[k + 1] - cp_array[k]

                if is_ansi:
                    parts.append(word_stream[fc: fc + char_count].decode('cp1252', errors='replace'))
                else:
                    parts.append(word_stream[fc: fc + char_count * 2].decode('utf-16-le', errors='replace'))

            ole.close()
            result = re.sub(r'[\x00-\x1f\x7f]', '', ''.join(parts))
            return result.strip()

        except Exception as e:
            logger.error(f"提取 doc 文本失败: {e}")
            return f"[doc 提取失败: {str(e)}]"

    @staticmethod
    async def _extract_xlsx_text(file_content: bytes) -> str:
        """Extract Excel text (supports .xlsx and legacy .xls)."""
        # xlsx (ZIP format)
        if file_content[:2] == b'PK':
            try:
                wb = openpyxl.load_workbook(io.BytesIO(file_content), read_only=True, data_only=True)
                parts = []
                for sheet in wb.worksheets:
                    parts.append(f"[Sheet: {sheet.title}]")
                    for row in sheet.iter_rows(values_only=True):
                        parts.append('\t'.join('' if v is None else str(v) for v in row))
                return '\n'.join(parts)
            except Exception as e:
                # openpyxl has zero tolerance for non-standard styles.xml (e.g. an empty
                # <fill/>) and raises TypeError: expected
                # <class 'openpyxl.styles.fills.Fill'>; calamine (a Rust implementation that
                # does not parse styles.xml) reads such files without error
                logger.warning(f"openpyxl 提取 xlsx 文本失败: {e}，尝试 calamine 降级读取")
                try:
                    from python_calamine import CalamineWorkbook
                    cwb = CalamineWorkbook.from_filelike(io.BytesIO(file_content))
                    parts = []
                    for sheet_name in cwb.sheet_names:
                        parts.append(f"[Sheet: {sheet_name}]")
                        for row in cwb.get_sheet_by_name(sheet_name).to_python():
                            parts.append('\t'.join('' if v is None else str(v) for v in row))
                    return '\n'.join(parts)
                except Exception as e_fallback:
                    logger.error(f"提取 xlsx 文本失败: openpyxl({e}), calamine({e_fallback})")
                    return f"[xlsx 提取失败: {str(e_fallback)}]"

        # xls (OLE2/BIFF format)
        try:
            import xlrd
            wb = xlrd.open_workbook(file_contents=file_content)
            parts = []
            for sheet in wb.sheets():
                parts.append(f"[Sheet: {sheet.name}]")
                for row_idx in range(sheet.nrows):
                    parts.append('\t'.join(str(sheet.cell_value(row_idx, col)) for col in range(sheet.ncols)))
            return '\n'.join(parts)
        except Exception as e:
            logger.error(f"提取 xls 文本失败: {e}")
            return f"[xls 提取失败: {str(e)}]"

    async def _extract_csv_text(self, file_content: bytes) -> str:
        """Extract CSV text."""
        try:
            text = self._decode_text_safe(file_content)
            reader = csv.reader(io.StringIO(text))
            return '\n'.join('\t'.join(row) for row in reader)
        except Exception as e:
            logger.error(f"提取 CSV 文本失败: {e}")
            return f"[CSV 提取失败: {str(e)}]"

    async def _extract_json_text(self, file_content: bytes) -> str:
        """Extract JSON text."""
        try:
            text = self._decode_text_safe(file_content)
            data = json.loads(text)
            return json.dumps(data, ensure_ascii=False, indent=2)
        except Exception as e:
            logger.error(f"提取 JSON 文本失败: {e}")
            return f"[JSON 提取失败: {str(e)}]"

    def _is_word_file(self, file_content: bytes, mime_type: str) -> bool:
        """Detect a Word file (doc / docx) without relying on the extension."""
        # Legacy .doc
        if mime_type == 'application/msword':
            return True

        # Modern .docx (the ZIP contains word/document.xml)
        header = file_content[:4]
        if header == b'PK\x03\x04':
            try:
                with zipfile.ZipFile(io.BytesIO(file_content)) as zf:
                    return "word/document.xml" in zf.namelist()
            except:
                pass

        return False

    def _is_excel_file(self, file_content: bytes, mime_type: str) -> bool:
        """Detect an Excel file (xls / xlsx) without relying on the extension."""
        # Legacy .xls
        if mime_type == 'application/vnd.ms-excel':
            return True

        # Modern .xlsx (the ZIP contains xl/workbook.xml)
        header = file_content[:4]
        if header == b'PK\x03\x04':
            try:
                with zipfile.ZipFile(io.BytesIO(file_content)) as zf:
                    return "xl/workbook.xml" in zf.namelist()
            except:
                pass

        return False

    @staticmethod
    def _decode_text_safe(file_content: bytes) -> str:
        """
        Universal text decoding.
        Auto-detects the encoding; utf-8 / gbk / gb2312 / utf-8-sig / ascii and more are supported.
        Never raises, never produces mojibake.
        """
        if not file_content:
            return ""

        # 1. Auto-detect the file encoding
        detect = chardet.detect(file_content)
        encoding = detect.get("encoding") or "utf-8"
        encoding = encoding.lower()

        # 2. Cover common Chinese encodings
        # Note: latin-1 must not be used as a fallback -- it maps bytes 1:1 and never fails,
        # "successfully" decoding binaries such as ZIP/PDF into strings containing NUL
        # (\\x00), which triggers CharacterNotInRepertoireError on write to PostgreSQL.
        compatible_encodings = ["utf-8", "gbk", "gb18030", "gb2312", "ascii"]

        # 3. Try decoders in priority order
        for enc in [encoding] + compatible_encodings:
            if not enc:
                continue
            try:
                return sanitize_text(file_content.decode(enc.strip()))
            except (UnicodeDecodeError, LookupError):
                continue

        # Last-resort fallback: replace invalid bytes with U+FFFD and strip NUL as well
        return sanitize_text(file_content.decode("utf-8", errors="replace"))


def get_multimodal_service(db: Session) -> MultimodalService:
    """Get a multimodal service instance (dependency injection)."""
    return MultimodalService(db)
